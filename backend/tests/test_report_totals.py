"""Report exports: Indian formatter + executive-summary totals math."""
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from app.services.report_service import build_report_pdf, compute_report_totals
from app.utils.money import format_inr

from tests.test_master_data import setup_refs
from tests.test_purchases import make_item
from tests.test_sales import sale_payload


# ---------- format_inr (the single INR formatter for all report output) ----------


def test_format_inr_indian_grouping():
    assert format_inr("123456.00") == "1,23,456.00"
    assert format_inr("1000") == "1,000.00"
    assert format_inr("999.99") == "999.99"
    assert format_inr("0") == "0.00"
    assert format_inr("10000000.10") == "1,00,00,000.10"
    assert format_inr("99999999.99") == "9,99,99,999.99"
    assert format_inr("-1234.56") == "-1,234.56"
    assert format_inr(Decimal("50")) == "50.00"
    assert format_inr(12345.5) == "12,345.50"
    assert format_inr("garbage") == "0.00"


# ---------- executive-summary totals (Decimal, reconciling) ----------


def _sale(subtotal: str, discount: str, total: str, method: str = "cash") -> dict:
    return {
        "subtotal": subtotal,
        "discount": discount,
        "total_amount": total,
        "payment_method": method,
    }


def test_compute_report_totals_reconcile():
    sales = [
        _sale("1000.00", "0.00", "1000.00", "cash"),
        _sale("500.00", "100.00", "400.00", "upi"),
        _sale("300.00", "0.00", "300.00", "card"),
    ]
    purchases = [{"total_amount": "200.00"}]
    expenses = [{"amount": "50.00"}]
    wastage = [{"estimated_value": "30.00"}]
    t = compute_report_totals(sales, purchases, expenses, wastage)

    assert t["gross"] == Decimal("1800.00")
    assert t["discounts"] == Decimal("100.00")
    assert t["net"] == Decimal("1700.00")
    assert t["outflow"] == Decimal("250.00")
    assert t["cash"] == Decimal("1450.00")
    assert t["adjusted"] == Decimal("1420.00")
    assert t["margin_pct"] == Decimal("1420.00") / Decimal("1700.00") * 100
    # Invariants: net = gross - discounts; cash = net - outflow;
    # adjusted = cash - wastage; the payment split sums back to net.
    assert t["net"] == t["gross"] - t["discounts"]
    assert t["cash"] == t["net"] - t["outflow"]
    assert t["adjusted"] == t["cash"] - t["wastage_memo"]
    assert sum(t["payment"].values()) == t["net"]
    assert t["payment"]["cash"] == Decimal("1000.00")
    assert t["payment"]["upi"] == Decimal("400.00")
    assert t["payment"]["card"] == Decimal("300.00")


def test_compute_report_totals_margin_na_without_sales():
    t = compute_report_totals([], [{"total_amount": "200.00"}], [], [])
    assert t["net"] == Decimal("0")
    assert t["margin_pct"] is None
    assert t["cash"] == Decimal("-200.00")
    assert t["payment"] == {}


def test_compute_report_totals_other_payment_methods_tracked():
    """Non cash/upi/card methods still reconcile the split to net sales."""
    sales = [_sale("100.00", "0.00", "100.00", "room_charge")]
    t = compute_report_totals(sales, [], [], [])
    assert t["payment"]["room_charge"] == Decimal("100.00")
    assert sum(t["payment"].values()) == t["net"]


# ---------- PDF smoke ----------


def _pdf_data(**over) -> dict:
    data = {
        "period": "custom",
        "range_start": datetime(2026, 10, 1, tzinfo=UTC),
        "range_end": datetime(2026, 10, 18, tzinfo=UTC),
        "business_unit_id": None,
        "bu_names": {},
        "sales": [],
        "purchases": [],
        "expenses": [],
        "wastage": [],
        "movements": [],
    }
    data.update(over)
    return data


def test_build_report_pdf_smoke_empty_period():
    pdf = build_report_pdf(_pdf_data())
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 500


def test_build_report_pdf_handles_long_items_and_other_methods():
    long_name = "Super Long Item Name That Wraps And Must Not Be Truncated " * 5
    sales = [{
        "sale_number": "SAL-0001",
        "sold_at": datetime(2026, 10, 10, 12, 0, tzinfo=UTC),
        "business_unit_id": "bu1",
        "items": [{"item_name": long_name, "quantity": 2}],
        "subtotal": "123456.00",
        "discount": "0.00",
        "total_amount": "123456.00",
        "payment_method": "room_charge",
    }]
    pdf = build_report_pdf(_pdf_data(
        sales=sales,
        bu_names={"bu1": "Restaurant"},
    ))
    assert pdf.startswith(b"%PDF")


# ---------- export endpoints: cancelled sales excluded ----------


async def test_export_excludes_cancelled_sales(client, owner_headers):
    """Cancelled orders (stock restored) must not appear in either export."""
    from openpyxl import load_workbook
    from tests.test_rooms import check_in_room, make_room

    bu_id, cat_id = await setup_refs(client, owner_headers)
    item = await make_item(
        client, owner_headers, bu_id, cat_id, name="Chai", sku="CHAI-1", opening=100
    )
    ok = await client.post(
        "/api/sales", json=sale_payload(bu_id, item["id"], key="rep-sale-ok"),
        headers=owner_headers,
    )
    assert ok.status_code == 201, ok.text

    # A cancellable order: room-charge sales stay PENDING until delivered.
    room = await make_room(client, owner_headers, bu_id, "REP-1")
    stay = await check_in_room(client, owner_headers, room["id"], "Report Guest")
    pending = await client.post(
        "/api/sales",
        json=sale_payload(
            bu_id, item["id"], key="rep-sale-cancel",
            payment_method="room_charge", stay_id=stay["id"],
        ),
        headers=owner_headers,
    )
    assert pending.status_code == 201, pending.text
    cancel = await client.delete(f"/api/sales/{pending.json()['id']}", headers=owner_headers)
    assert cancel.status_code == 200, cancel.text
    assert cancel.json()["order_status"] == "CANCELLED"

    params = {"period": "custom", "from": "2020-01-01", "to": "2100-01-01"}

    xlsx = await client.get("/api/reports/export", params=params, headers=owner_headers)
    assert xlsx.status_code == 200, xlsx.text
    wb = load_workbook(BytesIO(xlsx.content))
    ws = wb["Sales"]
    sale_numbers = [
        ws.cell(row=r, column=1).value for r in range(2, (ws.max_row or 1) + 1)
    ]
    assert ok.json()["sale_number"] in sale_numbers
    assert pending.json()["sale_number"] not in sale_numbers

    pdf = await client.get("/api/reports/export-pdf", params=params, headers=owner_headers)
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF")
